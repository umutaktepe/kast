"""Tests for CastTableWriter and docx table integration."""

import os
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn

from src.models import CastExtractionResult, CharacterStats
from src.table_writer import CastTableWriter, append_cast_table, save_result


def test_table_generation():
    """Verify basic cast table structure, headers, and character data."""
    doc = Document()
    writer = CastTableWriter()

    char1 = CharacterStats(name="K KOMŞU", first_seen_order=1)
    char1.add_line(page=1)
    char1.add_line(page=2)

    result = CastExtractionResult(characters=[char1], total_lines=2, total_pages=2)

    writer.append_cast_table(doc, result)

    # Document must have exactly 1 table
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 2  # 1 header + 1 data row

    # Header cells
    assert table.rows[0].cells[0].text == "Karakter"
    assert table.rows[0].cells[1].text == "Replik Sayısı"
    assert table.rows[0].cells[2].text == "Repliklerin Geçtiği Sayfalar"
    assert table.rows[0].cells[3].text == "Notlar"

    # Data row
    assert table.rows[1].cells[0].text == "K KOMŞU"
    assert table.rows[1].cells[1].text == "2"
    assert table.rows[1].cells[2].text == "1, 2"
    assert table.rows[1].cells[3].text == ""


def test_table_borders():
    """Verify that all cells have top, left, bottom, right borders set properly in XML."""
    doc = Document()
    writer = CastTableWriter()

    char1 = CharacterStats(name="PORORO", first_seen_order=1)
    char1.add_line(page=1)
    result = CastExtractionResult(characters=[char1], total_lines=1, total_pages=1)

    writer.append_cast_table(doc, result)

    table = doc.tables[0]
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = tcPr.find(qn("w:tcBorders"))
            assert tcBorders is not None, "w:tcBorders element missing on cell"
            for edge in ("top", "left", "bottom", "right"):
                edge_elem = tcBorders.find(qn(f"w:{edge}"))
                assert edge_elem is not None, f"Border {edge} missing"
                assert edge_elem.get(qn("w:val")) == "single"
                assert edge_elem.get(qn("w:sz")) == "4"
                assert edge_elem.get(qn("w:color")) == "000000"


def test_column_widths():
    """Verify that column widths match [2.0, 1.0, 2.3, 1.2] inches for all cells."""
    doc = Document()
    writer = CastTableWriter()

    char1 = CharacterStats(name="EDDY", first_seen_order=1)
    char1.add_line(page=1)
    result = CastExtractionResult(characters=[char1], total_lines=1, total_pages=1)

    writer.append_cast_table(doc, result)

    table = doc.tables[0]
    expected_widths = [Inches(2.0), Inches(1.0), Inches(2.3), Inches(1.2)]

    for row in table.rows:
        for idx, expected_width in enumerate(expected_widths):
            assert row.cells[idx].width == expected_width
            # Verify paragraph indents are zeroed to prevent inheriting document hanging indents
            p = row.cells[idx].paragraphs[0]
            assert p.paragraph_format.left_indent == Inches(0)
            assert p.paragraph_format.first_line_indent == Inches(0)


def test_header_bolding_and_styling():
    """Verify that header text is bold, 11pt, Arial and heading title is bold 14pt Arial."""
    doc = Document()
    writer = CastTableWriter()

    char1 = CharacterStats(name="LOOPY", first_seen_order=1)
    char1.add_line(page=1)
    result = CastExtractionResult(characters=[char1], total_lines=1, total_pages=1)

    writer.append_cast_table(doc, result)

    # Heading paragraph check
    heading_p = None
    for p in doc.paragraphs:
        if "KAST TABLOSU" in p.text:
            heading_p = p
            break
    assert heading_p is not None, "KAST TABLOSU title paragraph not found"
    assert heading_p.runs[0].bold is True
    assert heading_p.runs[0].font.size == Pt(14)
    assert heading_p.runs[0].font.name == "Arial"

    # Header row formatting
    table = doc.tables[0]
    # Header should not repeat across pages (w:tblHeader should be absent)
    assert table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader")) is None

    from docx.enum.table import WD_ALIGN_VERTICAL

    hdr_cells = table.rows[0].cells
    for cell in hdr_cells:
        assert cell.vertical_alignment == WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        assert len(p.runs) > 0
        assert p.runs[0].font.bold is True
        assert p.runs[0].font.size == Pt(11)
        assert p.runs[0].font.name == "Arial"

    # Data row formatting
    data_cells = table.rows[1].cells
    for cell in data_cells:
        assert cell.vertical_alignment == WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        assert len(p.runs) > 0
        assert p.runs[0].font.size == Pt(10)
        assert p.runs[0].font.name == "Arial"


def test_page_break_before_table():
    """Verify that a page break is inserted before the table when appending."""
    doc = Document()
    doc.add_paragraph("Existing script line before cast table")
    writer = CastTableWriter()

    char1 = CharacterStats(name="CRONG", first_seen_order=1)
    char1.add_line(page=3)
    result = CastExtractionResult(characters=[char1], total_lines=1, total_pages=3)

    writer.append_cast_table(doc, result)

    # Find page break element in paragraphs
    page_breaks = [
        child
        for p in doc.paragraphs
        for r in p.runs
        for child in r._r
        if child.tag.endswith("br") and child.get(qn("w:type")) == "page"
    ]
    assert len(page_breaks) >= 1, "Page break w:br[w:type='page'] not found in document"


def test_multiple_characters_with_notes():
    """Verify table with multiple characters and custom notes."""
    doc = Document()
    writer = CastTableWriter()

    char1 = CharacterStats(name="PORORO", first_seen_order=1, notes="Başrol")
    char1.add_line(page=1)
    char1.add_line(page=3)
    char1.add_line(page=5)

    char2 = CharacterStats(name="CRONG", first_seen_order=2, notes="Bebek dinozor")
    char2.add_line(page=2)

    result = CastExtractionResult(
        characters=[char1, char2],
        total_lines=4,
        total_pages=5
    )

    doc_returned = writer.append_cast_table(doc, result)
    assert doc_returned is doc  # Chainable return

    table = doc.tables[0]
    assert len(table.rows) == 3

    # Row 1 (Pororo)
    assert table.rows[1].cells[0].text == "PORORO"
    assert table.rows[1].cells[1].text == "3"
    assert table.rows[1].cells[2].text == "1, 3, 5"
    assert table.rows[1].cells[3].text == "Başrol"

    # Row 2 (Crong)
    assert table.rows[2].cells[0].text == "CRONG"
    assert table.rows[2].cells[1].text == "1"
    assert table.rows[2].cells[2].text == "2"
    assert table.rows[2].cells[3].text == "Bebek dinozor"


def test_empty_characters_result():
    """Verify that an empty character list generates an empty table with just headers."""
    doc = Document()
    writer = CastTableWriter()
    result = CastExtractionResult(characters=[], total_lines=0, total_pages=0)

    writer.append_cast_table(doc, result)

    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 1  # Only header row
    assert table.rows[0].cells[0].text == "Karakter"


def test_save_result_and_standalone_functions(tmp_path):
    """Verify standalone append_cast_table and save_result functions."""
    doc = Document()
    char1 = CharacterStats(name="POBY", first_seen_order=1)
    char1.add_line(page=1)
    result = CastExtractionResult(characters=[char1], total_lines=1, total_pages=1)

    append_cast_table(doc, result)

    out_file = str(tmp_path / "output_test.docx")
    saved_path = save_result(doc, out_file)
    assert os.path.exists(out_file)
    assert saved_path == out_file

    # Verify reloadable docx
    reloaded_doc = Document(out_file)
    assert len(reloaded_doc.tables) == 1
    assert reloaded_doc.tables[0].rows[1].cells[0].text == "POBY"


def test_append_cast_table_without_page_break():
    """Verify appending table without inserting a page break."""
    doc = Document()
    char1 = CharacterStats(name="POBY", first_seen_order=1)
    result = CastExtractionResult(characters=[char1], total_lines=1, total_pages=1)

    writer = CastTableWriter()
    writer.append_cast_table(doc, result, add_page_break=False)

    # If add_page_break is False, no page break paragraph was added
    assert len(doc.paragraphs) == 1  # Only the table heading paragraph
    assert doc.paragraphs[0].text == "KAST TABLOSU"
    assert len(doc.tables) == 1

