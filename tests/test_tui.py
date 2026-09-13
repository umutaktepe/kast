"""Tests for Textual TUI Application (src/tui.py)."""

import os
import pytest
from docx import Document
from textual.widgets import Button

from src.tui import KastApp


@pytest.mark.asyncio
async def test_tui_app_mount_and_widgets():
    """Verify KastApp loads widgets correctly in headless test."""
    app = KastApp()
    async with app.run_test() as pilot:
        assert "Kast" in app.title
        docx_input = app.query_one("#docx-path")
        assert docx_input is not None

        pdf_input = app.query_one("#pdf-path")
        assert pdf_input is not None

        sort_radios = app.query_one("#sort-radios")
        assert sort_radios is not None

        cb_inplace = app.query_one("#cb-inplace")
        cb_standalone = app.query_one("#cb-standalone")
        assert cb_inplace is not None
        assert cb_standalone is not None

        btn_extract = app.query_one("#btn-extract")
        assert btn_extract is not None


@pytest.mark.asyncio
async def test_tui_checkbox_mutual_exclusivity():
    """Verify checking in-place unchecks standalone and vice-versa."""
    app = KastApp()
    async with app.run_test() as pilot:
        cb_inplace = app.query_one("#cb-inplace")
        cb_standalone = app.query_one("#cb-standalone")

        # Check in-place
        cb_inplace.value = True
        await pilot.pause()
        assert cb_inplace.value is True

        # Now check standalone -> inplace should be unchecked
        cb_standalone.value = True
        await pilot.pause()
        assert cb_standalone.value is True
        assert cb_inplace.value is False

        # Check inplace again -> standalone should be unchecked
        cb_inplace.value = True
        await pilot.pause()
        assert cb_inplace.value is True
        assert cb_standalone.value is False


@pytest.mark.asyncio
async def test_tui_extract_empty_path_shows_error():
    """Verify clicking extract without path writes error message."""
    app = KastApp()
    async with app.run_test() as pilot:
        app.query_one("#btn-extract", Button).action_press()
        await pilot.pause()

        log_area = app.query_one("#log-area")
        lines = [line.text for line in log_area.lines]
        combined = " ".join(lines)
        assert "Lütfen geçerli bir .docx dosyası belirtin" in combined or "Hata" in combined


@pytest.mark.asyncio
async def test_tui_successful_extraction(tmp_path):
    """Verify successful extraction end-to-end through TUI."""
    # Create a small script docx
    test_docx = tmp_path / "sample_script.docx"
    doc = Document()
    doc.add_paragraph("FİLMİN ADI\tDENEME")
    doc.add_paragraph("ÇEVİRMEN\tTEST")
    doc.add_paragraph("00.10")
    doc.add_paragraph("ALICE\t- Merhaba dünya!")
    doc.add_paragraph("BOB\t- Selam Alice!")
    doc.save(str(test_docx))

    app = KastApp()
    async with app.run_test() as pilot:
        # Input path with enclosing quotes (simulating drag & drop)
        docx_input = app.query_one("#docx-path")
        docx_input.value = f"'{test_docx}'"

        app.query_one("#btn-extract", Button).action_press()
        await pilot.pause()

        expected_output = tmp_path / "sample_script_kast.docx"
        assert os.path.exists(expected_output)

        out_doc = Document(str(expected_output))
        assert len(out_doc.tables) == 1
        table = out_doc.tables[0]
        assert len(table.rows) == 3  # Header + ALICE + BOB


def test_clean_drag_drop_path():
    """Verify clean_drag_drop_path handles various terminal drag & drop formats."""
    from src.tui import clean_drag_drop_path

    # Standard path with quotes
    assert clean_drag_drop_path("'/home/user/film.docx'") == "/home/user/film.docx"
    assert clean_drag_drop_path('"/home/user/film.docx"') == "/home/user/film.docx"

    # file:// URI scheme
    assert clean_drag_drop_path("file:///home/user/film.docx") == "/home/user/film.docx"
    assert clean_drag_drop_path("'file:///home/user/film.docx'") == "/home/user/film.docx"

    # URL percent-encoding (spaces and non-ASCII)
    assert clean_drag_drop_path("file:///home/user/My%20Script.docx") == "/home/user/My Script.docx"
    assert clean_drag_drop_path("file:///home/user/T%C3%BCrk%C3%A7e.docx") == "/home/user/Türkçe.docx"

    # Whitespace and newlines
    assert clean_drag_drop_path("  /home/user/film.docx\r\n\t  ") == "/home/user/film.docx"
    assert clean_drag_drop_path("") == ""


@pytest.mark.asyncio
async def test_tui_drag_drop_paste_events():
    """Verify KastApp handles app-level drag-and-drop paste events."""
    from textual import events

    app = KastApp()
    async with app.run_test() as pilot:
        # 1. Drop a DOCX anywhere in the app
        pilot.app.post_message(events.Paste(text="file:///home/user/test%20drop.docx"))
        await pilot.pause()
        docx_val = app.query_one("#docx-path").value
        assert docx_val == "/home/user/test drop.docx"

        # 2. Drop a PDF anywhere in the app
        pilot.app.post_message(events.Paste(text="'file:///home/user/ref.pdf'"))
        await pilot.pause()
        pdf_val = app.query_one("#pdf-path").value
        assert pdf_val == "/home/user/ref.pdf"
        # docx-path must remain unchanged
        assert app.query_one("#docx-path").value == "/home/user/test drop.docx"


@pytest.mark.asyncio
async def test_tui_input_routing_swap():
    """Verify smart swap if user pastes PDF into DOCX input or vice versa."""
    app = KastApp()
    async with app.run_test() as pilot:
        docx_inp = app.query_one("#docx-path")
        pdf_inp = app.query_one("#pdf-path")

        # Paste PDF into DOCX input
        docx_inp.value = "file:///home/user/wrong_box.pdf"
        await pilot.pause()

        # Should be automatically moved to pdf_inp and cleared from docx_inp
        assert docx_inp.value == ""
        assert pdf_inp.value == "/home/user/wrong_box.pdf"


def test_select_file_dialog_cancelled(monkeypatch):
    """Verify select_file_dialog returns None immediately when cancelled."""
    import subprocess
    from src.tui import select_file_dialog

    class FakeProcess:
        returncode = 1
        stdout = ""

    call_count = 0

    def fake_run(*a, **kw):
        nonlocal call_count
        call_count += 1
        return FakeProcess()

    monkeypatch.setattr(subprocess, "run", fake_run)
    res = select_file_dialog("Test", ["docx"])
    assert res is None
    # Must only call zenity once, and NOT cascade to kdialog or tkinter!
    assert call_count == 1


@pytest.mark.asyncio
async def test_tui_sort_radio_options_visible():
    """Verify all 3 sorting options (including sort-name) are fully visible and not clipped."""
    app = KastApp()
    async with app.run_test(size=(100, 24)) as pilot:
        col = app.query_one(".options-col")
        b1 = app.query_one("#sort-appearance")
        b2 = app.query_one("#sort-count")
        b3 = app.query_one("#sort-name")

        # All 3 widgets exist and are RadioButtons
        assert b1 is not None
        assert b2 is not None
        assert b3 is not None

        # Verify b3 (sort-name) is strictly above the bottom border of options-col
        col_bottom_y = col.region.y + col.region.height - 1
        assert b3.region.y < col_bottom_y


@pytest.mark.asyncio
async def test_clean_radio_and_checkbox_rendering():
    """Verify CleanRadioButton and CleanCheckbox render ASCII safe indicators (*)/( ) and [X]/[ ]."""
    from src.tui import CleanRadioButton, CleanCheckbox

    app = KastApp()
    async with app.run_test() as pilot:
        r_app = app.query_one("#sort-appearance", CleanRadioButton)
        r_cnt = app.query_one("#sort-count", CleanRadioButton)
        cb_in = app.query_one("#cb-inplace", CleanCheckbox)
        cb_st = app.query_one("#cb-standalone", CleanCheckbox)

        # Initial state: sort-appearance is True, others False
        assert "(*)" in str(r_app.render())
        assert "( )" in str(r_cnt.render())
        assert "[ ]" in str(cb_in.render())
        assert "[ ]" in str(cb_st.render())

        # Toggle radio button
        r_cnt.value = True
        await pilot.pause()
        assert "( )" in str(r_app.render())
        assert "(*)" in str(r_cnt.render())

        # Toggle checkbox
        cb_in.value = True
        await pilot.pause()
        assert "[X]" in str(cb_in.render())
        assert "[ ]" in str(cb_st.render())


def test_select_file_dialog_windows_tkinter(monkeypatch, tmp_path):
    """Verify select_file_dialog uses Tkinter cleanly on Windows when available."""
    import platform
    import sys
    from unittest.mock import MagicMock
    from src.tui import select_file_dialog

    fake_file = tmp_path / "win_test.docx"
    fake_file.write_text("dummy")

    monkeypatch.setattr(platform, "system", lambda: "Windows")

    fake_tk = MagicMock()
    fake_filedialog = MagicMock()
    fake_tk.filedialog = fake_filedialog
    fake_filedialog.askopenfilename.return_value = str(fake_file)

    monkeypatch.setitem(sys.modules, "tkinter", fake_tk)
    monkeypatch.setitem(sys.modules, "tkinter.filedialog", fake_filedialog)

    selected = select_file_dialog("Test Windows", ["docx"])
    assert selected == os.path.normpath(str(fake_file))


def test_select_file_dialog_windows_powershell_fallback(monkeypatch, tmp_path):
    """Verify select_file_dialog falls back to PowerShell with CREATE_NO_WINDOW if tkinter fails."""
    import platform
    import subprocess
    from src.tui import select_file_dialog

    fake_file = tmp_path / "ps_test.docx"
    fake_file.write_text("dummy")

    monkeypatch.setattr(platform, "system", lambda: "Windows")

    # Simulate tkinter not working by raising in import or askopenfilename
    import sys
    monkeypatch.setitem(sys.modules, "tkinter", None)

    class FakePSProcess:
        returncode = 0
        stdout = str(fake_file) + "\n"

    passed_creationflags = []

    def fake_run(cmd, **kwargs):
        passed_creationflags.append(kwargs.get("creationflags"))
        return FakePSProcess()

    monkeypatch.setattr(subprocess, "run", fake_run)

    selected = select_file_dialog("Test PowerShell", ["docx"])
    assert selected == os.path.normpath(str(fake_file))
    assert 0x08000000 in passed_creationflags


def test_setup_windows_console_does_not_crash():
    """Verify setup_windows_console executes without error on any platform."""
    from src.tui import setup_windows_console
    setup_windows_console()



