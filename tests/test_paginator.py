"""Unit tests for the paginator engine."""

import os
from unittest.mock import MagicMock, patch
import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from src.models import DialogueLine
from src.parser import DubbingDocxParser, ParsedParagraph
from src.paginator import DocumentPaginator, Paginator, PurePythonLayoutPaginator


def test_layout_paginator_initialization():
    """Test initial dimensions and margins calculation."""
    paginator = PurePythonLayoutPaginator(
        page_height_pt=792,
        page_width_pt=612,
        top_margin_pt=72,
        bottom_margin_pt=72,
        left_margin_pt=72,
        right_margin_pt=72,
        dialogue_indent_pt=108,
    )
    assert paginator.printable_height == 648
    assert paginator.printable_width == 468
    assert paginator.dialogue_width == 360
    assert paginator.single_line_height == 16.5  # 11 * 1.5


def test_layout_paginator_advance_across_paragraphs():
    """Test that page numbers advance across paragraphs and don't remain stuck at 1."""
    paginator = PurePythonLayoutPaginator(
        page_height_pt=792,
        top_margin_pt=72,
        bottom_margin_pt=72,
    )
    # Printable height is 648 pt.
    # Single 1-line paragraph is (1 * 16.5) + 10 = 26.5 pt.
    # 30 paragraphs * 26.5 pt = 795 pt > 648 pt -> must wrap to page 2 and beyond!
    paragraphs = []
    for i in range(60):
        paragraphs.append(
            ParsedParagraph(
                index=i,
                text=f"SPEAKER\t- Line {i}",
                speaker="SPEAKER",
                dialogue=f"- Line {i}",
            )
        )

    assigned = paginator.paginate_paragraphs(paragraphs)
    assert len(assigned) == 60
    assert assigned[0].page == 1
    # Page must advance beyond 1
    assert assigned[-1].page >= 3
    # Check monotonicity
    pages = [p.page for p in assigned]
    assert sorted(pages) == pages
    # Check that both page 1 and page 2 are present
    assert 1 in pages
    assert 2 in pages
    assert 3 in pages


def test_layout_paginator_multiline_wrapping():
    """Test that long dialogue lines wrap and take more vertical height."""
    paginator = PurePythonLayoutPaginator()
    short_text = "- Hello"
    long_text = (
        "- Bu replik oldukça uzun bir cümledir ve standart mizanpajda "
        "kesinlikle birden fazla satıra sarılmak zorundadır. "
        "Arial 11 punto ve 1.5 satır aralığı ile kontrol ediyoruz."
    )
    short_lines = paginator.estimate_lines(short_text, paginator.dialogue_width)
    long_lines = paginator.estimate_lines(long_text, paginator.dialogue_width)
    assert short_lines == 1
    assert long_lines >= 3


def test_layout_paginator_fallback_without_font():
    """Test layout paginator works gracefully even if font loading fails."""
    paginator = PurePythonLayoutPaginator(font_path="/nonexistent/path/font.ttf")
    # Even with font=None, it should estimate lines and paginate properly
    p1 = ParsedParagraph(index=0, text="Header", is_metadata=True)
    assigned = paginator.paginate_paragraphs([p1])
    assert assigned[0].page == 1

    # 40 dialogue lines should advance beyond page 1 with fallback
    paras = [
        ParsedParagraph(
            index=i,
            text=f"CHAR\t- Fallback line {i} with some additional words for width",
            speaker="CHAR",
            dialogue=f"- Fallback line {i} with some additional words for width",
        )
        for i in range(40)
    ]
    assigned_many = paginator.paginate_paragraphs(paras)
    assert assigned_many[-1].page > 1


def test_document_paginator_default_process():
    """Test DocumentPaginator default tier (pure python layout)."""
    paginator = DocumentPaginator()
    p1 = ParsedParagraph(index=0, text="Header")
    p2 = ParsedParagraph(index=1, text="Dialogue 1", speaker="SAL", dialogue="- Test")
    assigned = paginator.process([p1, p2])
    assert len(assigned) == 2
    assert assigned[0].page == 1
    assert assigned[1].page == 1


def test_document_paginator_with_doc_dimensions():
    """Test DocumentPaginator extracts section dimensions from Document."""
    doc = Document()
    sec = doc.sections[0]
    from docx.shared import Pt
    sec.page_height = Pt(400)
    sec.top_margin = Pt(50)
    sec.bottom_margin = Pt(50)

    paginator = DocumentPaginator(doc)
    assert paginator.layout_paginator.printable_height == 300


def test_document_paginator_with_xml_hard_page_break():
    """Test DocumentPaginator detection of XML hard page breaks."""
    doc = Document()
    p0 = doc.add_paragraph("First page content")
    doc.add_page_break()  # empty paragraph with w:br w:type="page"
    p2 = doc.add_paragraph("Second page content")

    parsed_p0 = ParsedParagraph(index=0, text="First page content", speaker="A", dialogue="- Hi")
    parsed_p2 = ParsedParagraph(index=2, text="Second page content", speaker="B", dialogue="- Bye")

    paginator = DocumentPaginator(doc)
    assigned = paginator.process([parsed_p0, parsed_p2])
    assert assigned[0].page == 1
    assert assigned[1].page == 2


def test_document_paginator_with_xml_soft_page_break():
    """Test DocumentPaginator detection of Word lastRenderedPageBreak."""
    doc = Document()
    p0 = doc.add_paragraph("Paragraph 1")
    p1 = doc.add_paragraph()
    # Inject w:lastRenderedPageBreak into p1
    soft_break_xml = parse_xml(
        r'<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        r'<w:lastRenderedPageBreak/>'
        r'<w:t>Paragraph 2 on new page</w:t>'
        r'</w:r>'
    )
    p1._element.append(soft_break_xml)

    parsed_p0 = ParsedParagraph(index=0, text="Paragraph 1", speaker="A", dialogue="- P1")
    parsed_p1 = ParsedParagraph(index=1, text="Paragraph 2 on new page", speaker="B", dialogue="- P2")

    paginator = DocumentPaginator(doc)
    assigned = paginator.process([parsed_p0, parsed_p1])
    assert assigned[0].page == 1
    assert assigned[1].page == 2


def test_document_paginator_with_pdf():
    """Test DocumentPaginator Tier 1 PDF mapping with mock pdfplumber."""
    paragraphs = [
        ParsedParagraph(index=0, text="Header", is_metadata=True),
        ParsedParagraph(index=1, text="CHAR1\t- First line", speaker="CHAR1", dialogue="- First line"),
        ParsedParagraph(index=2, text="CHAR2\t- Second line", speaker="CHAR2", dialogue="- Second line"),
    ]

    mock_pdf = MagicMock()
    page1 = MagicMock()
    page1.extract_text.return_value = "Title Header\nCHAR1\tFirst line"
    page2 = MagicMock()
    page2.extract_text.return_value = "CHAR2\tSecond line\nEnd of scene"
    mock_pdf.pages = [page1, page2]
    mock_pdf.__enter__.return_value = mock_pdf

    with patch("pdfplumber.open", return_value=mock_pdf):
        with patch("os.path.exists", return_value=True):
            paginator = DocumentPaginator()
            assigned = paginator.process(paragraphs, pdf_path="dummy.pdf")
            assert assigned[0].page == 1
            assert assigned[1].page == 1
            assert assigned[2].page == 2


def test_paginator_interface_and_alias():
    """Test Paginator class alias and assign_pages classmethod."""
    p1 = ParsedParagraph(index=0, text="Test")
    assigned = Paginator.assign_pages(doc=None, parsed_paragraphs=[p1])
    assert len(assigned) == 1
    assert assigned[0].page == 1


def test_pagination_on_pororo_example():
    """Test pagination on actual example Pororo script."""
    example_path = "example/PORORO: SWEET CASTLE ADVENTURE.docx"
    if not os.path.exists(example_path):
        pytest.skip("Example file not found")

    doc = Document(example_path)
    parser = DubbingDocxParser()
    paragraphs = parser.parse_document_paragraphs(doc)
    assert len(paragraphs) > 1000

    paginator = DocumentPaginator(doc)
    assigned = paginator.process(paragraphs)

    # First dialogue line should be on page 1
    first_dialogue = next(p for p in assigned if p.speaker and p.dialogue)
    assert first_dialogue.page == 1

    # Total pages should be realistic (between 50 and 80 pages for ~1600 paragraphs)
    max_page = max(p.page for p in assigned)
    assert 50 <= max_page <= 80

    # Ensure page numbers are monotonically non-decreasing
    pages = [p.page for p in assigned]
    assert sorted(pages) == pages


def test_document_paginator_with_pdf_not_found():
    """Test DocumentPaginator raises FileNotFoundError when pdf doesn't exist."""
    paginator = DocumentPaginator()
    with pytest.raises(FileNotFoundError):
        paginator.process([], pdf_path="/nonexistent/path/file.pdf")


def test_layout_paginator_empty_paragraphs():
    """Test paginate_paragraphs handles empty paragraph list."""
    paginator = PurePythonLayoutPaginator()
    assert paginator.paginate_paragraphs([]) == []

