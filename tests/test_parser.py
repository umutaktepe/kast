import os
from docx import Document
from src.parser import DubbingDocxParser, ParsedParagraph
from src.models import DialogueLine


def test_detect_metadata_vs_dialogue():
    parser = DubbingDocxParser()

    # Metadata örneği (tire içermeyen veya bilinen header anahtarları)
    is_meta, key, val = parser.parse_header_line("FİLMİN ADI\tANOTHER END")
    assert is_meta is True
    assert key == "FİLMİN ADI"
    assert val == "ANOTHER END"

    # Case-insensitive header örneği
    is_meta, key, val = parser.parse_header_line("Filmin Adı\tANOTHER END")
    assert is_meta is True
    assert key == "Filmin Adı"
    assert val == "ANOTHER END"

    is_meta, key, val = parser.parse_header_line("filmin adı\tANOTHER END")
    assert is_meta is True
    assert key == "filmin adı"
    assert val == "ANOTHER END"

    # Diyalog örneği (sekme ve ardından tire ile başlayan satır)
    is_dial, char, text = parser.parse_dialogue_line("SAL\t- Hallettim.")
    assert is_dial is True
    assert char == "SAL"
    assert text == "- Hallettim."


def test_timecode_detection():
    parser = DubbingDocxParser()
    assert parser.is_timecode("01.59") is True
    assert parser.is_timecode("01.00.06") is True
    assert parser.is_timecode("01:23") is True
    assert parser.is_timecode("00.41") is True
    assert parser.is_timecode("01:00:06") is True
    assert parser.is_timecode("MC COOKIE") is False
    assert parser.is_timecode("") is False
    assert parser.is_timecode("123") is False
    assert parser.is_timecode("01:23:45:67") is True
    assert parser.is_timecode("01.23.45.67") is True
    assert parser.is_timecode("01:23:45:67:89") is False


def test_parse_header_line_edge_cases():
    parser = DubbingDocxParser()

    # Empty string
    assert parser.parse_header_line("") == (False, None, None)
    assert parser.parse_header_line("   ") == (False, None, None)

    # No tab separator
    assert parser.parse_header_line("JUST SOME TEXT") == (False, None, None)

    # Dialogue line shouldn't be parsed as header
    assert parser.parse_header_line("SAL\t- Hallettim.") == (False, None, None)

    # Known metadata keys
    is_meta, key, val = parser.parse_header_line("ÇEVİRMEN\tUMUT AKTEPE")
    assert is_meta is True
    assert key == "ÇEVİRMEN"
    assert val == "UMUT AKTEPE"

    is_meta, key, val = parser.parse_header_line("STÜDYO\tAK'S DUBBING")
    assert is_meta is True
    assert key == "STÜDYO"
    assert val == "AK'S DUBBING"


def test_parse_dialogue_line_edge_cases():
    parser = DubbingDocxParser()

    # Empty string
    assert parser.parse_dialogue_line("") == (False, None, None)

    # No tab
    assert parser.parse_dialogue_line("SAL - Hallettim.") == (False, None, None)

    # Different dash characters (en-dash, em-dash, standard hyphen)
    assert parser.parse_dialogue_line("SAL\t– En dash replik")[0] is True
    assert parser.parse_dialogue_line("SAL\t— Em dash replik")[0] is True

    # Header line shouldn't be parsed as dialogue
    assert parser.parse_dialogue_line("FİLMİN ADI\tPORORO")[0] is False

    # Speaker name cleaning (colon, whitespace, trailing punctuation)
    is_dial, char, text = parser.parse_dialogue_line("  PORORO: \t- Merhaba! ")
    assert is_dial is True
    assert char == "PORORO"
    assert text == "- Merhaba!"


def test_parse_document_paragraphs_mock():
    parser = DubbingDocxParser()
    doc = Document()
    doc.add_paragraph("FİLMİN ADI\tTEST MOVIE")
    doc.add_paragraph("ÇEVİRMEN\tTEST ÇEVİRMEN")
    doc.add_paragraph("")
    doc.add_paragraph("01.00")
    doc.add_paragraph("CHAR A\t- İlk replik.")
    doc.add_paragraph("01.05")
    doc.add_paragraph("CHAR B\t- İkinci replik.")
    doc.add_paragraph("Açıklama notu")

    parsed = parser.parse_document_paragraphs(doc)
    assert len(parsed) == 7  # 2 headers + 2 timecodes + 2 dialogues + 1 plain text (empty skipped)

    # Header 1
    assert parsed[0].is_metadata is True
    assert parsed[0].meta_key == "FİLMİN ADI"
    assert parsed[0].meta_value == "TEST MOVIE"

    # Header 2
    assert parsed[1].is_metadata is True
    assert parsed[1].meta_key == "ÇEVİRMEN"
    assert parsed[1].meta_value == "TEST ÇEVİRMEN"

    # Timecode 1
    assert parsed[2].is_timecode is True
    assert parsed[2].timecode == "01.00"

    # Dialogue 1
    assert parsed[3].speaker == "CHAR A"
    assert parsed[3].dialogue == "- İlk replik."
    assert parsed[3].timecode == "01.00"

    # Timecode 2
    assert parsed[4].is_timecode is True
    assert parsed[4].timecode == "01.05"

    # Dialogue 2
    assert parsed[5].speaker == "CHAR B"
    assert parsed[5].dialogue == "- İkinci replik."
    assert parsed[5].timecode == "01.05"

    # Plain text paragraph
    assert parsed[6].text == "Açıklama notu"
    assert parsed[6].speaker is None
    assert parsed[6].dialogue is None
    assert parsed[6].is_metadata is False
    assert parsed[6].is_timecode is False


def test_parsed_paragraph_to_dialogue_line():
    p_dial = ParsedParagraph(
        index=3,
        text="SAL\t- Selam",
        speaker="SAL",
        dialogue="- Selam",
        timecode="01.20",
        page=2,
    )
    d_line = p_dial.to_dialogue_line()
    assert isinstance(d_line, DialogueLine)
    assert d_line.speaker == "SAL"
    assert d_line.text == "- Selam"
    assert d_line.timecode == "01.20"
    assert d_line.page == 2
    assert d_line.paragraph_index == 3

    p_meta = ParsedParagraph(
        index=0,
        text="FİLMİN ADI\tTEST",
        is_metadata=True,
        meta_key="FİLMİN ADI",
        meta_value="TEST",
    )
    assert p_meta.to_dialogue_line() is None


def test_pororo_example_integration():
    example_path = "example/PORORO: SWEET CASTLE ADVENTURE.docx"
    assert os.path.exists(example_path)

    parser = DubbingDocxParser()
    doc = Document(example_path)
    parsed = parser.parse_document_paragraphs(doc)

    dialogue_paragraphs = [p for p in parsed if p.speaker and p.dialogue]
    metadata_paragraphs = [p for p in parsed if p.is_metadata]

    # Exactly 982 dialogue lines as per specification
    assert len(dialogue_paragraphs) == 982

    # Exactly 2 metadata headers (FİLMİN ADI, ÇEVİRMEN)
    assert len(metadata_paragraphs) == 2
    meta_dict = {p.meta_key: p.meta_value for p in metadata_paragraphs}
    assert meta_dict.get("FİLMİN ADI") == "PORORO: SWEET CASTLE ADVENTURE"
    assert meta_dict.get("ÇEVİRMEN") == "UMUT AKTEPE"

    # No metadata mistaken for dialogue
    speaker_names = {p.speaker for p in dialogue_paragraphs}
    for meta_key in parser.known_metadata_keys:
        assert meta_key not in speaker_names
    assert "FİLMİN ADI" not in speaker_names
    assert "ÇEVİRMEN" not in speaker_names


def test_parse_docx_and_parse_paragraphs():
    example_path = "example/PORORO: SWEET CASTLE ADVENTURE.docx"
    parser = DubbingDocxParser()

    # Test parse_docx
    parsed_paras = parser.parse_docx(example_path)
    assert len(parsed_paras) == 1630

    # Test parse_paragraphs interface
    meta, dial_lines = parser.parse_paragraphs(example_path)
    assert meta.get("FİLMİN ADI") == "PORORO: SWEET CASTLE ADVENTURE"
    assert meta.get("ÇEVİRMEN") == "UMUT AKTEPE"
    assert len(dial_lines) == 982
    assert all(isinstance(dl, DialogueLine) for dl in dial_lines)
    assert dial_lines[0].speaker == "MC COOKIE"
    assert dial_lines[0].timecode == "00.41"


def test_header_heuristics_and_custom_keys():
    # Heuristic matching: uppercase key containing AD, ÇEVİR, TARİH, METİN, PROJE, KOD
    parser = DubbingDocxParser()
    is_meta, key, val = parser.parse_header_line("PROJE KODU\tKST-001")
    assert is_meta is True
    assert key == "PROJE KODU"
    assert val == "KST-001"

    # Too many words (>4) should not match heuristic
    is_meta, _, _ = parser.parse_header_line("BU COK UZUN BIR PROJE KODU BASLIGI\tDEGER")
    assert is_meta is False

    # Custom known keys
    custom_parser = DubbingDocxParser(known_metadata_keys={"OZEL_BASLIK"})
    is_meta, key, val = custom_parser.parse_header_line("OZEL_BASLIK\tDEGER")
    assert is_meta is True
    assert key == "OZEL_BASLIK"

