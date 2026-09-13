"""Integration tests for dubbing cast extraction pipeline (Task 5)."""

import os
import shutil
import pytest
from docx import Document

from kast import process_cast_document, main, build_cli_parser


def test_full_pipeline_on_example(tmp_path):
    """Verify end-to-end pipeline on example PORORO document."""
    src_file = "example/PORORO: SWEET CASTLE ADVENTURE.docx"
    assert os.path.exists(src_file)

    target_copy = tmp_path / "test_pororo.docx"
    shutil.copy(src_file, target_copy)

    output_path = process_cast_document(str(target_copy))
    assert os.path.exists(output_path)
    assert output_path.endswith("_kast.docx")

    # Çıktı dosyasını doğrula
    doc = Document(output_path)
    assert len(doc.tables) >= 1
    table = doc.tables[-1]

    # Karakter sayısını doğrula (Pororo'da en az 30 karakter)
    assert len(table.rows) > 30

    # Karakter isimlerini kontrol et
    char_names = [row.cells[0].text for row in table.rows[1:]]
    assert "MC COOKIE" in char_names
    assert "PORORO" in char_names
    assert "SUGAR QUEEN" in char_names

    # Tanımlayıcıların karaktere dönüşmediğini doğrula
    assert "FİLMİN ADI" not in char_names
    assert "ÇEVİRMEN" not in char_names

    # Varsayılan sıralama ilk görünme sırasıdır (appearance)
    assert char_names[0] == "MC COOKIE"


def test_pipeline_sort_by_count(tmp_path):
    """Verify sorting by line count descending."""
    src_file = "example/PORORO: SWEET CASTLE ADVENTURE.docx"
    target_copy = tmp_path / "test_pororo_count.docx"
    shutil.copy(src_file, target_copy)

    output_path = process_cast_document(str(target_copy), sort_by="count")
    doc = Document(output_path)
    table = doc.tables[-1]

    counts = [int(row.cells[1].text) for row in table.rows[1:]]
    assert len(counts) > 30
    # Ensure descending order
    assert counts == sorted(counts, reverse=True)
    # Top character has highest lines
    top_char = table.rows[1].cells[0].text
    assert top_char == "PORORO"


def test_pipeline_sort_by_name(tmp_path):
    """Verify sorting alphabetically by character name."""
    src_file = "example/PORORO: SWEET CASTLE ADVENTURE.docx"
    target_copy = tmp_path / "test_pororo_name.docx"
    shutil.copy(src_file, target_copy)

    output_path = process_cast_document(str(target_copy), sort_by="name")
    doc = Document(output_path)
    table = doc.tables[-1]

    char_names = [row.cells[0].text for row in table.rows[1:]]
    assert len(char_names) > 30
    assert char_names == sorted(char_names)


def test_pipeline_custom_output_path(tmp_path):
    """Verify custom output_path."""
    src_file = "example/PORORO: SWEET CASTLE ADVENTURE.docx"
    target_copy = tmp_path / "test_pororo_custom.docx"
    custom_out = tmp_path / "custom_dir" / "my_cast_output.docx"
    os.makedirs(custom_out.parent, exist_ok=True)
    shutil.copy(src_file, target_copy)

    returned_path = process_cast_document(str(target_copy), output_path=str(custom_out))
    assert returned_path == str(custom_out)
    assert os.path.exists(custom_out)


def test_pipeline_in_place(tmp_path):
    """Verify in-place update by passing target_copy as output_path."""
    src_file = "example/PORORO: SWEET CASTLE ADVENTURE.docx"
    target_copy = tmp_path / "test_pororo_inplace.docx"
    shutil.copy(src_file, target_copy)

    returned_path = process_cast_document(str(target_copy), output_path=str(target_copy))
    assert returned_path == str(target_copy)

    doc = Document(str(target_copy))
    assert len(doc.tables) >= 1
    assert len(doc.tables[-1].rows) > 30


def test_pipeline_standalone(tmp_path):
    """Verify standalone cast document contains only cast table."""
    src_file = "example/PORORO: SWEET CASTLE ADVENTURE.docx"
    target_copy = tmp_path / "test_pororo_standalone.docx"
    shutil.copy(src_file, target_copy)

    output_path = process_cast_document(str(target_copy), standalone=True)
    assert os.path.exists(output_path)

    doc = Document(output_path)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) > 30

    # Standalone document shouldn't contain the 900+ dialogue paragraphs
    # Only heading paragraph(s) for the table
    assert len(doc.paragraphs) <= 5


def test_pipeline_file_not_found():
    """Verify FileNotFoundError is raised if document does not exist."""
    with pytest.raises(FileNotFoundError):
        process_cast_document("non_existent_file_xyz_123.docx")


def test_cli_argument_parsing(tmp_path):
    """Verify CLI parser options."""
    parser = build_cli_parser()
    args = parser.parse_args([
        "sample.docx",
        "-o", "out.docx",
        "--sort", "count",
        "--in-place",
        "--pdf", "sample.pdf",
        "--standalone"
    ])
    assert args.docx_file == "sample.docx"
    assert args.output_path == "out.docx"
    assert args.sort_by == "count"
    assert args.in_place is True
    assert args.pdf_path == "sample.pdf"
    assert args.standalone is True


def test_cli_main_execution(tmp_path):
    """Verify main() CLI execution with argv."""
    src_file = "example/PORORO: SWEET CASTLE ADVENTURE.docx"
    target_copy = tmp_path / "cli_pororo.docx"
    shutil.copy(src_file, target_copy)

    ret = main([str(target_copy), "--sort", "appearance"])
    assert ret == 0

    expected_out = str(tmp_path / "cli_pororo_kast.docx")
    assert os.path.exists(expected_out)


def test_cli_main_interactive(tmp_path, monkeypatch):
    """Verify main() interactive prompt when docx_file not passed via argv."""
    src_file = "example/PORORO: SWEET CASTLE ADVENTURE.docx"
    target_copy = tmp_path / "interactive_pororo.docx"
    shutil.copy(src_file, target_copy)

    # Simulate drag & drop with enclosing quotes
    quoted_input = f"'{target_copy}'"
    monkeypatch.setattr("builtins.input", lambda _: quoted_input)

    ret = main([])
    assert ret == 0

    expected_out = str(tmp_path / "interactive_pororo_kast.docx")
    assert os.path.exists(expected_out)


def test_cli_main_conflict(tmp_path, capsys):
    """Verify error when both --in-place and --standalone are provided."""
    ret = main(["dummy.docx", "--in-place", "--standalone"])
    assert ret == 1
    captured = capsys.readouterr()
    assert "birlikte kullanılamaz" in captured.out or "birlikte kullanılamaz" in captured.err


def test_cli_main_empty_input(monkeypatch):
    """Verify error when empty input is provided."""
    monkeypatch.setattr("builtins.input", lambda _: "")
    ret = main([])
    assert ret == 1
