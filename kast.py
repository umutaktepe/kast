import re
import sys
import pandas as pd
from collections import defaultdict
import pdfplumber
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def extract_cast_from_pdf(pdf_path):
    # Karakterlerin adlarını, replik sayısını ve geçtiği sayfaları tutmak için sözlük
    character_data = defaultdict(lambda: {"count": 0, "pages": set()})

    with pdfplumber.open(pdf_path) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            text = page.extract_text()
            if not text:
                print(f"Sayfa {page_number} boş veya metin çıkarılamadı.")
                continue

            # Geçersiz başlıkları ve süre bilgilerini hariç tutan regex
            filtered_lines = [
                line for line in text.splitlines()
                if not re.match(r"^(F\u0130LM ADI|\u00c7EV\u0130REN|\d{2}\.\d{2})", line)
            ]

            # Karakter adı ve replikleri tespit etmek için regex (noktalama ve boşluk içeren tam isimleri almak için güncellendi)
            matches = re.findall(r"^([A-Z\u00c7\u011e\u0130\u00d6\u015e\u00dc][A-Z\u00c7\u011e\u0130\u00d6\u015e\u00dc0-9\s\.\-]*)\s*-\s+", "\n".join(filtered_lines), re.MULTILINE)

            if not matches:
                print(f"Sayfa {page_number} için eşleşme bulunamadı.")
                continue

            for character in matches:
                character = character.strip().rstrip("-")  # Gereksiz tire işaretlerini kaldır
                character_data[character]["count"] += 1
                character_data[character]["pages"].add(page_number)

    # Sonuçları tabloya dönüştürme
    data = []
    for character, info in character_data.items():
        data.append({
            "Karakter": character,
            "Replik Sayısı": info["count"],
            "Repliklerin Geçtiği Sayfalar": ", ".join(map(str, sorted(info["pages"]))),
            "Notlar": ""
        })

    if not data:
        print("PDF'den hiçbir veri çıkarılamadı. PDF yapısını kontrol edin.")
    else:
        print(f"Çıkarılan veriler: {data}")

    df = pd.DataFrame(data)
    return df

def set_table_borders(table):
    # Tabloya kenarlık ekler
    tbl = table._element
    for cell in tbl.xpath(".//w:tc"):  # Her hücre için
        tc_pr = cell.xpath(".//w:tcPr")[0]
        borders = OxmlElement("w:tcBorders")
        for border_name in ["top", "left", "bottom", "right"]:
            border = OxmlElement(f"w:{border_name}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4")  # Kalınlık
            border.set(qn("w:space"), "0")
            border.set(qn("w:color"), "000000")
            borders.append(border)
        tc_pr.append(borders)

def save_to_docx(df, output_path):
    document = Document()

    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"

    # Başlık satırı
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Karakter"
    hdr_cells[1].text = "Replik Sayısı"
    hdr_cells[2].text = "Repliklerin Geçtiği Sayfalar"
    hdr_cells[3].text = "Notlar"

    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(12)  # Yazı boyutunu büyüt

    # Verileri doldurma
    for _, row in df.iterrows():
        row_cells = table.add_row().cells
        row_cells[0].text = row["Karakter"]
        row_cells[1].text = str(row["Replik Sayısı"])
        row_cells[2].text = row["Repliklerin Geçtiği Sayfalar"]
        row_cells[3].text = row["Notlar"]

    set_table_borders(table)  # Kenarlıkları ayarla
    document.save(output_path)
    print(f"Sonuçlar {output_path} dosyasına kaydedildi.")

if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Lütfen kast için PDF dosyasının yolunu terminale sürükleyip bırakın.")
        input_file = input("Dosya yolunu buraya yapıştırın: ").strip()
    else:
        input_file = sys.argv[1].strip("'\"")  # Tırnak işaretlerini temizle

    # Dosya yolundaki fazladan tırnak işaretlerini temizle
    if input_file.startswith("'") and input_file.endswith("'"):
        input_file = input_file[1:-1]
    elif input_file.startswith('"') and input_file.endswith('"'):
        input_file = input_file[1:-1]

    docx_output_path = "kast.docx"  # Çıktı DOCX dosyasının adı

    df = extract_cast_from_pdf(input_file)

    # Sonuçları DOCX dosyasına kaydetme
    if not df.empty:
        save_to_docx(df, docx_output_path)
    else:
        print("Hiçbir veri bulunamadı. PDF yapısını kontrol edin.")
