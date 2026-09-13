"""DOCX Dubbing format parser engine for Kast."""

import re
from dataclasses import dataclass
from typing import Dict, List, Optional, Set, Tuple, Union
from docx import Document

from src.models import DialogueLine

KNOWN_METADATA_KEYS: Set[str] = {
    "FİLMİN ADI",
    "FİLM ADI",
    "DİZİ ADI",
    "DİZİ",
    "ÇEVİRMEN",
    "ÇEVİREN",
    "SEZON",
    "BÖLÜM",
    "BÖLÜM ADI",
    "KAYIT",
    "STÜDYO",
    "YÖNETMEN",
    "SESLENDİRME YÖNETMENİ",
    "TARİH",
    "TITLE",
    "TRANSLATOR",
}

TIMECODE_REGEX = re.compile(r"^\d{2}[\.:]\d{2}(?:[\.:]\d{2})?$")


@dataclass
class ParsedParagraph:
    """Represents an analyzed paragraph from a dubbing script DOCX."""

    index: int
    text: str
    speaker: Optional[str] = None
    dialogue: Optional[str] = None
    timecode: Optional[str] = None
    is_metadata: bool = False
    is_timecode: bool = False
    meta_key: Optional[str] = None
    meta_value: Optional[str] = None
    page: int = 1

    def to_dialogue_line(self) -> Optional[DialogueLine]:
        """Convert to DialogueLine model if this paragraph is a dialogue."""
        if self.speaker and self.dialogue:
            return DialogueLine(
                speaker=self.speaker,
                text=self.dialogue,
                timecode=self.timecode,
                page=self.page,
                paragraph_index=self.index,
            )
        return None


class DubbingDocxParser:
    """Parser engine for dubbing scripts formatted in DOCX."""

    def __init__(self, known_metadata_keys: Optional[Set[str]] = None) -> None:
        self.known_metadata_keys: Set[str] = (
            set(known_metadata_keys) if known_metadata_keys else set(KNOWN_METADATA_KEYS)
        )

    def is_timecode(self, text: str) -> bool:
        """Check if paragraph text represents a timecode (e.g. 00.41, 01.59, 01.00.06)."""
        cleaned = text.strip()
        return bool(TIMECODE_REGEX.match(cleaned))

    def parse_header_line(self, text: str) -> Tuple[bool, Optional[str], Optional[str]]:
        """
        Identify if text is a header metadata line (e.g. FİLMİN ADI <TAB> TITLE).
        Returns (is_metadata, key, value).
        """
        text_clean = text.strip()
        if not text_clean:
            return False, None, None

        parts = text_clean.split("\t", 1)
        if len(parts) == 2:
            key, val = parts[0].strip(), parts[1].strip()
            if key in self.known_metadata_keys or any(key.startswith(k) for k in self.known_metadata_keys):
                return True, key, val
            if not val.startswith("-") and not val.startswith("–") and not val.startswith("—") and key.isupper() and len(key.split()) <= 4:
                # Muhtemel başlık tanımlayıcısı
                if any(m in key for m in ["AD", "ÇEVİR", "TARİH", "METİN", "PROJE", "KOD"]):
                    return True, key, val
        return False, None, None

    def parse_dialogue_line(self, text: str) -> Tuple[bool, Optional[str], Optional[str]]:
        """
        Identify if text is a dialogue line (SPEAKER <TAB> - Replik).
        Returns (is_dialogue, speaker, dialogue_text).
        """
        text_clean = text.strip()
        if not text_clean:
            return False, None, None

        parts = text_clean.split("\t", 1)
        if len(parts) == 2:
            speaker, dialogue = parts[0].strip(), parts[1].strip()
            # Standart format: Diyalog tire ile başlar (- Replik)
            if dialogue.startswith("-") or dialogue.startswith("–") or dialogue.startswith("—"):
                # Karakter ismi temizliği (büyük harf ve varsa gereksiz işaretler)
                speaker_clean = speaker.strip(" \t:.-")
                return True, speaker_clean, dialogue
        return False, None, None

    def parse_document_paragraphs(self, doc: Document) -> List[ParsedParagraph]:
        """Parse all paragraphs of a docx Document into ParsedParagraph structures."""
        parsed: List[ParsedParagraph] = []
        current_timecode: Optional[str] = None

        for idx, p in enumerate(doc.paragraphs):
            text = p.text.strip()
            if not text:
                continue

            # 1. Süre kodu kontrolü
            if self.is_timecode(text):
                current_timecode = text
                parsed.append(ParsedParagraph(index=idx, text=text, is_timecode=True, timecode=text))
                continue

            # 2. Başlık / Metadata kontrolü
            is_meta, m_key, m_val = self.parse_header_line(p.text)
            if is_meta:
                parsed.append(
                    ParsedParagraph(
                        index=idx,
                        text=text,
                        is_metadata=True,
                        meta_key=m_key,
                        meta_value=m_val,
                    )
                )
                continue

            # 3. Diyalog kontrolü
            is_dial, speaker, dial_text = self.parse_dialogue_line(p.text)
            if is_dial:
                parsed.append(
                    ParsedParagraph(
                        index=idx,
                        text=text,
                        speaker=speaker,
                        dialogue=dial_text,
                        timecode=current_timecode,
                    )
                )
            else:
                # Bilinmeyen / düz metin paragrafı
                parsed.append(ParsedParagraph(index=idx, text=text))

        return parsed

    def parse_docx(self, docx_path: str) -> List[ParsedParagraph]:
        """Convenience helper to parse directly from a file path."""
        doc = Document(docx_path)
        return self.parse_document_paragraphs(doc)

    def parse_paragraphs(
        self, docx_path_or_doc: Union[str, Document]
    ) -> Tuple[Dict[str, str], List[DialogueLine]]:
        """
        Parses document and returns (metadata_dict, dialogue_lines_list).
        Produces interface compatible with the task brief.
        """
        if isinstance(docx_path_or_doc, str):
            doc = Document(docx_path_or_doc)
        else:
            doc = docx_path_or_doc

        parsed = self.parse_document_paragraphs(doc)
        metadata: Dict[str, str] = {}
        dialogue_lines: List[DialogueLine] = []

        for p in parsed:
            if p.is_metadata and p.meta_key:
                metadata[p.meta_key] = p.meta_value or ""
            elif p.speaker and p.dialogue:
                d_line = p.to_dialogue_line()
                if d_line:
                    dialogue_lines.append(d_line)

        return metadata, dialogue_lines
