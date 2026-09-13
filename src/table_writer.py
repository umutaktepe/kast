"""DOCX Table Writer for Dubbing Cast Extraction (Kast 2.0).

Generates and formats the 4-column cast table and appends it directly to
a Word (.docx) document with proper styling, borders, and column widths.
"""
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from src.models import CastExtractionResult


def set_cell_border(cell, **kwargs) -> None:
    """Apply w:tcBorders XML element to a table cell.

    Args:
        cell: python-docx table cell object.
        **kwargs: Edge specifications e.g. top={"sz": 4, "val": "single", "color": "000000"}.
    """
    tcPr = cell._tc.get_or_add_tcPr()
    existing_borders = tcPr.find(qn("w:tcBorders"))
    if existing_borders is not None:
        tcPr.remove(existing_borders)

    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f"w:{edge}"
            element = OxmlElement(tag)
            element.set(qn("w:val"), str(edge_data.get("val", "single")))
            element.set(qn("w:sz"), str(edge_data.get("sz", 4)))
            element.set(qn("w:space"), str(edge_data.get("space", 0)))
            element.set(qn("w:color"), str(edge_data.get("color", "000000")))
            tcBorders.append(element)
    tcPr.append(tcBorders)


class CastTableWriter:
    """Formats and writes dubbing cast statistics table to DOCX documents."""

    def __init__(self):
        pass

    def append_cast_table(
        self,
        doc: Document,
        result: CastExtractionResult,
        add_page_break: bool = True
    ) -> Document:
        """Append a formatted 4-column cast table to the given Word document.

        Args:
            doc: Document instance to append the table to.
            result: CastExtractionResult containing character statistics.
            add_page_break: Whether to insert a page break before the table.

        Returns:
            The Document instance for method chaining.
        """
        # 1. Metnin sonuna yeni sayfa kesmesi ekle
        if add_page_break:
            doc.add_page_break()

        # 2. Tablo başlığı
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.left_indent = Inches(0)
        p.paragraph_format.right_indent = Inches(0)
        p.paragraph_format.first_line_indent = Inches(0)
        run = p.add_run("KAST TABLOSU")
        run.bold = True
        run.font.size = Pt(14)
        run.font.name = "Arial"

        # 3. 4 sütunlu tablo oluştur
        table = doc.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Sütun genişlikleri (Karakter, Replik Sayısı, Sayfalar, Notlar)
        col_widths = [Inches(2.0), Inches(1.0), Inches(2.3), Inches(1.2)]
        for idx, width in enumerate(col_widths):
            table.columns[idx].width = width

        # Kenarlık özellikleri (İnce siyah tek çizgi)
        border_spec = {"val": "single", "sz": 4, "color": "000000"}
        all_borders = {k: border_spec for k in ["top", "left", "bottom", "right"]}

        # 4. Başlık satırı
        hdr_row = table.rows[0]
        hdr_trPr = hdr_row._tr.get_or_add_trPr()
        hdr_trPr.append(OxmlElement("w:cantSplit"))

        hdr_cells = hdr_row.cells
        headers = ["Karakter", "Replik Sayısı", "Repliklerin Geçtiği Sayfalar", "Notlar"]

        for i, text in enumerate(headers):
            hdr_cells[i].text = text
            hdr_cells[i].width = col_widths[i]
            set_cell_border(hdr_cells[i], **all_borders)
            hp = hdr_cells[i].paragraphs[0]
            hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
            hp.paragraph_format.space_before = Pt(2)
            hp.paragraph_format.space_after = Pt(2)
            hp.paragraph_format.left_indent = Inches(0)
            hp.paragraph_format.right_indent = Inches(0)
            hp.paragraph_format.first_line_indent = Inches(0)
            if hp.runs:
                hp.runs[0].font.bold = True
                hp.runs[0].font.name = "Arial"
                hp.runs[0].font.size = Pt(11)

        # 5. Veri satırları
        for char in result.characters:
            row = table.add_row()
            row_trPr = row._tr.get_or_add_trPr()
            row_trPr.append(OxmlElement("w:cantSplit"))

            row_cells = row.cells
            row_cells[0].text = char.name or ""
            row_cells[1].text = str(char.line_count)
            row_cells[2].text = char.formatted_pages
            row_cells[3].text = char.notes or ""

            for i in range(4):
                row_cells[i].width = col_widths[i]
                set_cell_border(row_cells[i], **all_borders)
                rp = row_cells[i].paragraphs[0]
                rp.paragraph_format.space_before = Pt(2)
                rp.paragraph_format.space_after = Pt(2)
                rp.paragraph_format.left_indent = Inches(0)
                rp.paragraph_format.right_indent = Inches(0)
                rp.paragraph_format.first_line_indent = Inches(0)
                if rp.runs:
                    rp.runs[0].font.name = "Arial"
                    rp.runs[0].font.size = Pt(10)

        return doc

    def save_result(self, doc: Document, output_path: str) -> str:
        """Save the document to the specified output file path.

        Args:
            doc: Document instance to save.
            output_path: Target .docx file path.

        Returns:
            The saved output path.
        """
        doc.save(output_path)
        return output_path


def append_cast_table(
    doc: Document,
    result: CastExtractionResult,
    add_page_break: bool = True
) -> Document:
    """Convenience function to append cast table using default CastTableWriter."""
    writer = CastTableWriter()
    return writer.append_cast_table(doc, result, add_page_break=add_page_break)


def save_result(doc: Document, output_path: str) -> str:
    """Convenience function to save document to file path."""
    writer = CastTableWriter()
    return writer.save_result(doc, output_path)
